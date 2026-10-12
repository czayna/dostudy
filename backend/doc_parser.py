import os
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

def extract_text_from_file(file_path: Path, mime_type: str = "") -> str:
    ext = file_path.suffix.lower()
    text = ""
    try:
        if ext == ".pdf":
            reader = PdfReader(str(file_path))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        elif ext in [".docx", ".doc"]:
            doc = Document(str(file_path))
            for p in doc.paragraphs:
                if p.text:
                    text += p.text + "\n"
        elif ext in [".pptx", ".ppt"]:
            prs = Presentation(str(file_path))
            for slide_idx, slide in enumerate(prs.slides, 1):
                text += f"\n--- Slide {slide_idx} ---\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text += shape.text + "\n"
        elif ext in [".txt", ".md"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
    except Exception as e:
        print(f"Error parsing {file_path}: {e}")
    return text.strip()
